import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement

BASE_DIR = r"d:\Machinelearning icfoss\FINAL PROJECT FOR GRADING"
SCREENSHOT_DIR = os.path.join(BASE_DIR, "Screenshot of  the website")
SC1 = os.path.join(SCREENSHOT_DIR, "Screenshot 1 .jpeg")
SC2 = os.path.join(SCREENSHOT_DIR, "screenshot 2.jpeg")

def get_plot(filename):
    return os.path.join(BASE_DIR, "plots", filename)

def set_cell_background(cell, color):
    shading_elm_1 = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
    cell._tc.get_or_add_tcPr().append(shading_elm_1)

def add_heading(doc, text, level=1):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 0 else WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(31, 78, 121) # Deep blue
    run.font.bold = True
    run.font.size = Pt(16) if level == 0 else Pt(14)
    return heading

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Page ")
    add_page_number(p.add_run())

# --- PAGE 1: COVER ---
doc.add_paragraph('\n' * 5)
title = add_heading(doc, 'PROJECT REPORT', level=0)
title.runs[0].font.size = Pt(28)
doc.add_paragraph('\n' * 2)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Flight Fare Prediction using Machine Learning\n\n")
r.font.size = Pt(16)
r.font.bold = True
r.font.name = 'Times New Roman'

det = doc.add_paragraph()
det.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = det.add_run("Submitted in partial fulfillment of the requirements for the\nHybrid Internship Program on Machine Learning through Python\n\n\n\n")
r2.font.size = Pt(14)
r2.font.name = 'Times New Roman'

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = sub2.add_run("Submitted by:\nKARTHIK KS\n\n\n\n")
r3.font.size = Pt(14)
r3.font.bold = True
r3.font.name = 'Times New Roman'

inst = doc.add_paragraph()
inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = inst.add_run("International Centre for Free and Open Source Software (ICFOSS)\nGovernment of Kerala\nJune 2026")
r4.font.size = Pt(14)
r4.font.bold = True
r4.font.name = 'Times New Roman'
doc.add_page_break()

# --- PAGE 2: ABSTRACT ---
add_heading(doc, 'ABSTRACT', level=0)
doc.add_paragraph("This report presents the learning experience of a participant in the training / internship programme conducted at ICFOSS. The programme introduced participants to practical computing concepts, open-source tools, software development practices, and problem-solving methods relevant to modern technology projects.")
doc.add_paragraph("The training followed a structured approach consisting of orientation, theoretical sessions, hands-on exercises, demonstrations, mini-project activity, review, and documentation. Participants were encouraged to understand the need for open-source technology, apply tools in practical situations, and prepare a concise report of their learning.")
doc.add_paragraph("The report covers the organization profile, objectives of the programme, methodology adopted, tools and technologies introduced, learning outcomes, challenges faced, and future scope. It has been customized to detail the completion of the Flight Fare Prediction project.")
doc.add_paragraph("As part of the project activity, a machine learning pipeline was developed in Python to predict domestic flight ticket prices. Using a dataset of 300,153 records, basic data exploration and preprocessing were carried out. Five regression algorithms were tested, and the Random Forest Regressor was found to be the best model, achieving an R² score of 98.56%. Finally, the trained model was successfully deployed as a web application using Flask.")
doc.add_page_break()

# --- PAGE 3: TABLE OF CONTENTS ---
add_heading(doc, 'TABLE OF CONTENTS', level=0)
toc_table = doc.add_table(rows=1, cols=2)
toc_table.style = 'Table Grid'
hdr_cells = toc_table.rows[0].cells
set_cell_background(hdr_cells[0], 'DAE6F3')
set_cell_background(hdr_cells[1], 'DAE6F3')
hdr_cells[0].paragraphs[0].add_run('Section').font.bold = True
hdr_cells[1].paragraphs[0].add_run('Page No.').font.bold = True
hdr_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

sections = [
    ("Abstract", "2"),
    ("Chapter 1 - Introduction", "4"),
    ("  Organization Profile", "4"),
    ("  Programme Background", "4"),
    ("  Objectives of the Programme", "4"),
    ("Chapter 2 - Methodology and Activities", "5"),
    ("  Training Methodology", "5"),
    ("  Tools and Technologies Used", "5"),
    ("  Activity Flow", "5"),
    ("Chapter 3 - Learning Outcomes and Results", "7"),
    ("  Skills Developed", "7"),
    ("  Participant Outcomes", "7"),
    ("  Findings", "7"),
    ("Summary, Challenges, Future Scope and Conclusion", "9"),
    ("References and Annexure", "10")
]
for sec, page in sections:
    row_cells = toc_table.add_row().cells
    row_cells[0].text = sec
    row_cells[1].text = page
    row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
doc.add_page_break()

# --- CHAPTER 1 ---
add_heading(doc, 'CHAPTER 1 - INTRODUCTION', level=0)
add_heading(doc, 'Organization Profile', level=1)
doc.add_paragraph("The International Centre for Free and Open Source Software (ICFOSS) is an autonomous organization established by the Government of Kerala to promote Free and Open Source Software, research, localization, capacity building, and technology adoption for public benefit.")

add_heading(doc, 'Programme Background', level=1)
doc.add_paragraph("The programme was designed to introduce participants to practical computing concepts and open-source development practices. It focused on connecting classroom knowledge with real-world implementation through guided sessions and hands-on activities.")

add_heading(doc, 'Objectives of the Programme', level=1)
doc.add_paragraph("To understand the role of open-source tools in education, research, and software development.", style='List Bullet')
doc.add_paragraph("To gain practical exposure to selected tools, workflows, and development practices.", style='List Bullet')
doc.add_paragraph("To improve problem-solving, documentation, teamwork, and presentation skills.", style='List Bullet')
doc.add_paragraph("To complete a small activity or mini-project and summarize the learning outcomes.", style='List Bullet')
doc.add_page_break()

# --- CHAPTER 2 ---
add_heading(doc, 'CHAPTER 2 - METHODOLOGY AND ACTIVITIES', level=0)
add_heading(doc, 'Training Methodology', level=1)
doc.add_paragraph("The training methodology combined concept explanation, live demonstration, guided practice, independent exploration, review, and documentation. Each topic was introduced with its purpose, followed by implementation steps and participant interaction.")

add_heading(doc, 'Tools and Technologies Used', level=1)
doc.add_paragraph("Operating system and command-line basics for working with files and tools.", style='List Bullet')
doc.add_paragraph("Programming environment for writing, running, and debugging simple programs (e.g., Python 3.13, Flask).", style='List Bullet')
doc.add_paragraph("Open-source libraries or platforms selected according to the programme topic (e.g., Pandas, Scikit-learn, Matplotlib).", style='List Bullet')
doc.add_paragraph("Documentation tools for preparing notes, reports, and final submission files.", style='List Bullet')

add_heading(doc, 'Activity Flow', level=1)
act_table = doc.add_table(rows=1, cols=6)
act_table.style = 'Table Grid'
hdr_cells = act_table.rows[0].cells
for i, text in enumerate(["Orientation", "Concepts", "Hands-on", "Mini Task", "Review", "Report"]):
    set_cell_background(hdr_cells[i], 'DAE6F3')
    hdr_cells[i].paragraphs[0].add_run(text).font.bold = True
    hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)

doc.add_paragraph("\nThe internship progressed systematically, beginning with an orientation to open-source concepts, followed by theoretical concept learning in Data Science. Participants then engaged in hands-on implementation of machine learning models. This culminated in full-scale project development, where models were trained and deployed. Finally, the work underwent peer and mentor review, leading to the formal documentation and final submission of the project.")

doc.add_paragraph("Project Specific Methodology:", style='Heading 2').runs[0].font.color.rgb = RGBColor(31, 78, 121)

doc.add_paragraph("Data Loading & Exploration: The dataset containing 300,153 flight booking records was loaded using Pandas. Initial checks were performed to understand the structure of the data.", style='List Number')

if os.path.exists(get_plot('08_airline_count_and_price.png')):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(get_plot('08_airline_count_and_price.png'), width=Inches(3.6))
    p_cap = doc.add_paragraph("Figure 1: Airline Count and Average Price Distribution")
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].font.italic = True

doc.add_paragraph("Exploratory Data Analysis (EDA): Graphs were created using Matplotlib and Seaborn to visualize the data. It was observed that flight duration and days left until departure strongly affect the ticket price.", style='List Number')

if os.path.exists(get_plot('05_days_left_vs_price.png')):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(get_plot('05_days_left_vs_price.png'), width=Inches(3.6))
    p_cap = doc.add_paragraph("Figure 2: Impact of Days Left on Ticket Price")
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].font.italic = True

doc.add_paragraph("Data Preprocessing: Outliers in the dataset were removed using the IQR method. Categorical data (like airlines and cities) was converted into numerical format using One-Hot Encoding and Label Encoding so the machine learning models could process it.", style='List Number')

doc.add_paragraph("Model Training: The cleaned data was split into training and testing sets. Five different regression models were trained on the data and evaluated to see which one performed the best.", style='List Number')

doc.add_paragraph("Web Application Deployment (Flask): The best performing model was saved and integrated into a simple web application using Flask, allowing users to input flight details and see the predicted price.", style='List Number')

doc.add_page_break()

# --- CHAPTER 3 ---
add_heading(doc, 'CHAPTER 3 - LEARNING OUTCOMES AND RESULTS', level=0)
add_heading(doc, 'Skills Developed', level=1)
doc.add_paragraph("Technical understanding of the selected topic and its practical use.", style='List Bullet')
doc.add_paragraph("Ability to follow structured workflows for implementation and testing.", style='List Bullet')
doc.add_paragraph("Basic debugging and error-analysis skills during hands-on practice.", style='List Bullet')
doc.add_paragraph("Improved communication through discussion, documentation, and presentation.", style='List Bullet')

add_heading(doc, 'Participant Outcomes', level=1)
doc.add_paragraph("By the end of the programme, the participant was able to explain the main concepts introduced during the sessions, identify the tools used, complete the assigned practical activity, and document the process in a clear report format.")

add_heading(doc, 'Findings', level=1)
doc.add_paragraph("The programme showed that guided hands-on learning improves technical confidence. Participants were able to connect theoretical ideas with implementation steps when each concept was supported by examples, demonstrations, and feedback.")

doc.add_paragraph("\nProject Specific Findings:", style='Heading 2').runs[0].font.color.rgb = RGBColor(31, 78, 121)
res_table = doc.add_table(rows=1, cols=4)
res_table.style = 'Table Grid'
hdr_cells = res_table.rows[0].cells
for i, text in enumerate(["Algorithm", "MAE (INR)", "RMSE (INR)", "R² Score"]):
    set_cell_background(hdr_cells[i], 'DAE6F3')
    hdr_cells[i].paragraphs[0].add_run(text).font.bold = True

results = [
    ("Random Forest", "1,250", "2,707", "0.9856"),
    ("XGBoost", "1,599", "2,880", "0.9837"),
    ("Decision Tree", "1,372", "3,159", "0.9804"),
    ("Linear Regression", "4,476", "6,731", "0.9112"),
    ("Support Vector Machine", "7,419", "12,514", "0.6931")
]
for row in results:
    row_cells = res_table.add_row().cells
    for i in range(4): row_cells[i].text = row[i]

doc.add_paragraph("\nThe Random Forest algorithm provided the best results, achieving an R² score of 98.56%. It also had the lowest errors, with a Mean Absolute Error (MAE) of INR 1,250. Random Forest performed better than standard linear models because it could handle the complex patterns in the airline pricing data more effectively.")

if os.path.exists(get_plot('12_model_comparison_r2.png')):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(get_plot('12_model_comparison_r2.png'), width=Inches(3.6))
    p_cap = doc.add_paragraph("Figure 3: R² Score Comparison Across Models")
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].font.italic = True

if os.path.exists(get_plot('15_feature_importance_rf.png')):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(get_plot('15_feature_importance_rf.png'), width=Inches(3.6))
    p_cap = doc.add_paragraph("Figure 4: Feature Importance (Random Forest)")
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].font.italic = True

doc.add_page_break()

# --- CONCLUSION ---
add_heading(doc, 'SUMMARY, CHALLENGES AND CONCLUSION', level=0)
add_heading(doc, 'Challenges', level=1)
doc.add_paragraph("Training machine learning models on a large dataset of over 300,000 records took a significant amount of time and computer memory.", style='List Bullet')
doc.add_paragraph("Converting text data like city names and airlines into numbers resulted in many new columns, making the data slightly harder to manage.", style='List Bullet')
doc.add_paragraph("Connecting the trained Python model to the Flask web page so that it could predict prices smoothly without errors.", style='List Bullet')

add_heading(doc, 'Scope for Future Learning', level=1)
doc.add_paragraph("In the future, the project could be improved by connecting it to live flight data APIs instead of using a static CSV file. The web application could also be hosted on the cloud (like AWS or Google Cloud) so anyone can access it on the internet. Finally, trying out deep learning techniques might improve the prediction accuracy even more.")

add_heading(doc, 'Conclusion', level=1)
doc.add_paragraph("The training / internship programme provided useful exposure to open-source technology, practical implementation, and professional documentation. The sessions helped the participant develop technical awareness, confidence, and a foundation for further study or project work.")

doc.add_page_break()

# --- REFERENCES AND ANNEXURE ---
add_heading(doc, 'REFERENCES AND ANNEXURE', level=0)

doc.add_paragraph("References", style='Heading 2').runs[0].font.color.rgb = RGBColor(31, 78, 121)
doc.add_paragraph("ICFOSS official resources and training materials.", style='List Bullet')
doc.add_paragraph("Session notes, demonstrations, and participant activity records.", style='List Bullet')
doc.add_paragraph("Open-source documentation related to the tools used in the programme (Python, Scikit-learn, Flask).", style='List Bullet')
doc.add_paragraph("Flight Price Prediction Dataset from Kaggle.", style='List Bullet')

doc.add_paragraph("\nAnnexure A: Web Application Screenshots", style='Heading 2').runs[0].font.color.rgb = RGBColor(31, 78, 121)
if os.path.exists(SC1):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run().add_picture(SC1, width=Inches(4.5))
    p1_cap = doc.add_paragraph("Figure 5: Flask Web Application - Input Interface")
    p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1_cap.runs[0].font.italic = True

if os.path.exists(SC2):
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run().add_picture(SC2, width=Inches(4.5))
    p2_cap = doc.add_paragraph("Figure 6: Flask Web Application - Prediction Result")
    p2_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_cap.runs[0].font.italic = True

doc.add_paragraph("\nAnnexure B: Core Code Implementation", style='Heading 2').runs[0].font.color.rgb = RGBColor(31, 78, 121)

doc.add_paragraph("Snippet 1: Data Preprocessing (Encoding & Scaling)")
code1 = doc.add_paragraph()
r_code1 = code1.add_run(
    "from sklearn.preprocessing import StandardScaler\n"
    "import pandas as pd\n\n"
    "# Apply One-Hot Encoding to categorical nominal features\n"
    "df = pd.get_dummies(df, columns=['airline', 'source_city', 'destination_city'])\n\n"
    "# Scale numerical features\n"
    "scaler = StandardScaler()\n"
    "df[['duration', 'days_left']] = scaler.fit_transform(df[['duration', 'days_left']])\n"
)
r_code1.font.name = 'Courier New'
r_code1.font.size = Pt(9)

doc.add_paragraph("Snippet 2: Model Training (Random Forest)")
code2 = doc.add_paragraph()
r_code2 = code2.add_run(
    "from sklearn.ensemble import RandomForestRegressor\n"
    "from sklearn.metrics import mean_absolute_error, r2_score\n\n"
    "# Train Random Forest Regressor\n"
    "rf_model = RandomForestRegressor(n_estimators=100, random_state=42, n_jobs=-1)\n"
    "rf_model.fit(X_train, y_train)\n\n"
    "# Evaluate model performance\n"
    "y_pred = rf_model.predict(X_test)\n"
    "print(f'R2 Score: {r2_score(y_test, y_pred)}')\n"
)
r_code2.font.name = 'Courier New'
r_code2.font.size = Pt(9)

doc.add_paragraph("Snippet 3: Flask Route for Prediction Inference")
code3 = doc.add_paragraph()
r_code3 = code3.add_run(
    "@app.route('/predict', methods=['POST'])\n"
    "def predict():\n"
    "    features = process_input(request.form) # Custom preprocessing function\n"
    "    predicted_price = model.predict(features)[0]\n"
    "    return jsonify({'predicted_price': round(predicted_price, 2)})\n"
)
r_code3.font.name = 'Courier New'
r_code3.font.size = Pt(9)

try:
    doc.save(os.path.join(BASE_DIR, 'ICFOSS_Project_Report_Final_Internship.docx'))
    print("Document generated successfully as Final Internship!")
except PermissionError:
    doc.save(os.path.join(BASE_DIR, 'ICFOSS_Project_Report_Final_Internship_v2.docx'))
    print("Document generated successfully as Final Internship v2!")
