import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

BASE_DIR = r"d:\Machinelearning icfoss\FINAL PROJECT FOR GRADING"
PLOT_DIR = os.path.join(BASE_DIR, "plots")
SC1 = os.path.join(BASE_DIR, "Screenshot of  the website", "Screenshot 1 .jpeg")

prs = Presentation()

def add_title_slide(title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = title_text
    subtitle.text = subtitle_text
    return slide

def add_content_slide(title_text, bullet_points):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = title_text
    tf = body.text_frame
    tf.text = bullet_points[0]
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
    return slide

def add_image_slide(title_text, img_path):
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title only layout
    title = slide.shapes.title
    title.text = title_text
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1), Inches(2), width=Inches(8))
    return slide

# Slide 1: Title
add_title_slide("Flight Fare Prediction", "Using Machine Learning\n\nSubmitted by: Karthik KS\nICFOSS Hybrid Internship Program")

# Slide 2: Program Background
add_content_slide("Program Background", [
    "Conducted by International Centre for Free and Open Source Software (ICFOSS).",
    "Focuses on practical computing and open-source development.",
    "Bridges the gap between theory and hands-on implementation.",
    "Final objective: Develop and deploy a full machine learning model."
])

# Slide 3: Problem Statement
add_content_slide("Problem Statement", [
    "Airline ticket prices are highly dynamic and fluctuate based on various factors.",
    "Goal: Build a predictive model to estimate domestic flight prices in India.",
    "Provides users with a clear expectation of ticket costs before booking.",
    "Allows airlines to understand which features drive pricing the most."
])

# Slide 4: Dataset Used
add_content_slide("The Dataset", [
    "Source: Kaggle (Flight Price Prediction Dataset).",
    "Size: Over 300,153 flight booking records.",
    "Key Features Used:",
    " - Airline Name (Vistara, Air India, Indigo, etc.)",
    " - Source & Destination Cities",
    " - Days Left until Departure",
    " - Flight Duration"
])

# Slide 5: Methodology
add_content_slide("Methodology & Preprocessing", [
    "1. Data Cleaning: Removed empty rows and invalid data.",
    "2. Outlier Removal: Used the IQR (Interquartile Range) method to remove crazy price spikes.",
    "3. Encoding: Used One-Hot Encoding to convert text (like Airline names) into numbers (0s and 1s).",
    "4. Scaling: Standardized numerical data (like duration) so the model trains fairly.",
    "5. Splitting: 80% data for training, 20% for testing."
])

# Slide 6: Visualizations
add_image_slide("Exploratory Data Analysis", os.path.join(PLOT_DIR, "08_airline_count_and_price.png"))

# Slide 7: Model Comparison
add_image_slide("Comparing Machine Learning Models", os.path.join(PLOT_DIR, "12_model_comparison_r2.png"))

# Slide 8: The Best Model
add_content_slide("Results: Why Random Forest?", [
    "Random Forest Regressor was the absolute best performer.",
    "R² Score: 98.56% (It explains 98.5% of the variance in ticket prices).",
    "Mean Absolute Error (MAE): Only ~1,250 INR.",
    "Why? Random Forest is an 'ensemble' model (it combines many decision trees), making it perfect for complex, non-linear airline pricing data."
])

# Slide 9: Flask Web App
add_image_slide("Web Application Deployment (Flask)", SC1)

# Slide 10: Conclusion
add_title_slide("Live Demo & Q&A", "Thank You!")

prs.save(os.path.join(BASE_DIR, "ICFOSS_Project_Presentation.pptx"))
print("Presentation generated successfully!")
